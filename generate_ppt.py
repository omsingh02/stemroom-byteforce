#!/usr/bin/env python3
"""
Generate Review 1 Presentation for Stemroom (Problem Statement 22: Online Music Collaboration Platform)
Team: ByteForce (Om Singh, Aryan Sharma, Adi)
Follows guidelines from docs/REVIEW-1.md and docs/DESIGN.md
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Design System Palette (from docs/DESIGN.md)
BG_COLOR = RGBColor(0xE5, 0xE8, 0xE2)       # Pale grey-green canvas (#e5e8e2)
PANEL_COLOR = RGBColor(0xF1, 0xF3, 0xEE)    # Panel/Card background (#f1f3ee)
INK = RGBColor(0x1C, 0x22, 0x20)            # Deep green-black text (#1c2220)
INK_MUTED = RGBColor(0x4F, 0x5A, 0x54)      # Secondary text / labels (#4f5a54)
LINE_COLOR = RGBColor(0xBC, 0xC5, 0xBD)     # Subtle divider lines (#bcc5bd)
ACCENT_MUSTARD = RGBColor(0xE0, 0xA5, 0x26) # Signature mustard accent (#e0a526)
COLOR_OK = RGBColor(0x2C, 0x6A, 0x43)       # Forest green / Approved (#2c6a43)
COLOR_DANGER = RGBColor(0xA8, 0x36, 0x2A)   # Deep red / Alert (#a8362a)
COLOR_WAIT = RGBColor(0x7D, 0x5A, 0x00)     # Amber / Pending (#7d5a00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT_HEADING = "Arial"
FONT_BODY = "Verdana"

def set_slide_background(slide):
    # Create full-slide background rectangle
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = BG_COLOR
    bg_shape.line.fill.background()
    return bg_shape

def add_header(slide, title_text, category="REVIEW 1 : STEMROOM"):
    # Category Tracker
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = category.upper()
    p_c.font.name = FONT_HEADING
    p_c.font.size = Pt(11)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_WAIT

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.name = FONT_HEADING
    p_t.font.size = Pt(26)
    p_t.font.bold = True
    p_t.font.color.rgb = INK

    # Subtle bottom rule
    rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.02)
    )
    rule.fill.solid()
    rule.fill.fore_color.rgb = LINE_COLOR
    rule.line.fill.background()

def create_card(slide, left, top, width, height, bg=PANEL_COLOR, border=LINE_COLOR):
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg
    if border:
        card.line.color.rgb = border
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Accent decorative block
    acc_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(0.2), Inches(4.8))
    acc_bar.fill.solid()
    acc_bar.fill.fore_color.rgb = ACCENT_MUSTARD
    acc_bar.line.fill.background()

    # Main Project & Subtitle Box
    t_box = s1.shapes.add_textbox(Inches(1.3), Inches(1.2), Inches(7.2), Inches(3.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "PROBLEM STATEMENT 22"
    p0.font.name = FONT_HEADING
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = INK_MUTED

    p1 = tf1.add_paragraph()
    p1.text = "Stemroom"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(46)
    p1.font.bold = True
    p1.font.color.rgb = INK

    p2 = tf1.add_paragraph()
    p2.text = "Online Music Collaboration Platform"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(20)
    p2.font.color.rgb = COLOR_OK

    p3 = tf1.add_paragraph()
    p3.text = "Review 1: Architecture, Core Java, JDBC & UI Implementation"
    p3.font.name = FONT_BODY
    p3.font.size = Pt(14)
    p3.font.color.rgb = INK_MUTED

    # Team Card
    create_card(s1, Inches(8.8), Inches(1.5), Inches(3.7), Inches(4.5), bg=PANEL_COLOR, border=LINE_COLOR)
    team_box = s1.shapes.add_textbox(Inches(9.1), Inches(1.8), Inches(3.1), Inches(3.9))
    tf_team = team_box.text_frame
    tf_team.word_wrap = True

    tp0 = tf_team.paragraphs[0]
    tp0.text = "TEAM DETAILS"
    tp0.font.name = FONT_HEADING
    tp0.font.size = Pt(12)
    tp0.font.bold = True
    tp0.font.color.rgb = COLOR_WAIT

    tp1 = tf_team.add_paragraph()
    tp1.text = "ByteForce"
    tp1.font.name = FONT_HEADING
    tp1.font.size = Pt(24)
    tp1.font.bold = True
    tp1.font.color.rgb = INK

    tp2 = tf_team.add_paragraph()
    tp2.text = "\nTeam Members:"
    tp2.font.name = FONT_HEADING
    tp2.font.size = Pt(13)
    tp2.font.bold = True
    tp2.font.color.rgb = INK_MUTED

    members = ["• Om Singh", "• Aryan Sharma", "• Adi"]
    for m in members:
        pm = tf_team.add_paragraph()
        pm.text = m
        pm.font.name = FONT_BODY
        pm.font.size = Pt(15)
        pm.font.color.rgb = INK

    # Banner image if present
    banner_path = "docs/brand/banner.png"
    if os.path.exists(banner_path):
        try:
            s1.shapes.add_picture(banner_path, Inches(1.3), Inches(4.7), width=Inches(7.0))
        except Exception:
            pass

    # =========================================================================
    # SLIDE 2: The Problem in One Sentence & Target Users
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "1. Problem Understanding & Motivation")

    # Core Problem Quote Card
    create_card(s2, Inches(0.8), Inches(1.8), Inches(11.733), Inches(1.6), bg=PANEL_COLOR, border=ACCENT_MUSTARD)
    prob_box = s2.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(11.1), Inches(1.4))
    tf_prob = prob_box.text_frame
    tf_prob.word_wrap = True

    pp0 = tf_prob.paragraphs[0]
    pp0.text = "THE CORE PROBLEM STATEMENT"
    pp0.font.name = FONT_HEADING
    pp0.font.size = Pt(11)
    pp0.font.bold = True
    pp0.font.color.rgb = COLOR_WAIT

    pp1 = tf_prob.add_paragraph()
    pp1.text = '"Musicians collaborating remotely lack a dedicated, structured hub for their audio stems, portfolio, and project milestones—leaving files scattered across chat apps with no clarity on versions or deadlines."'
    pp1.font.name = FONT_BODY
    pp1.font.size = Pt(16)
    pp1.font.italic = True
    pp1.font.color.rgb = INK

    # Two Column Breakdown
    # Col 1: Who Has This Problem
    create_card(s2, Inches(0.8), Inches(3.7), Inches(5.7), Inches(3.2))
    b1_box = s2.shapes.add_textbox(Inches(1.0), Inches(3.8), Inches(5.3), Inches(2.9))
    tf_b1 = b1_box.text_frame
    tf_b1.word_wrap = True
    p = tf_b1.paragraphs[0]
    p.text = "Target Audience & User Pain Points"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = INK

    points_1 = [
        "Independent musicians, songwriters, and audio producers.",
        "Remote collaborators across different geographic zones.",
        "Audio stems and mix takes get buried in informal messaging apps.",
        "Confusion regarding master versions and pending deliverables.",
        "Lack of an integrated portfolio to showcase vetted original work."
    ]
    for pt in points_1:
        p_pt = tf_b1.add_paragraph()
        p_pt.text = "• " + pt
        p_pt.font.name = FONT_BODY
        p_pt.font.size = Pt(13)
        p_pt.font.color.rgb = INK_MUTED

    # Col 2: The Solution: Stemroom
    create_card(s2, Inches(6.8), Inches(3.7), Inches(5.7), Inches(3.2))
    b2_box = s2.shapes.add_textbox(Inches(7.0), Inches(3.8), Inches(5.3), Inches(2.9))
    tf_b2 = b2_box.text_frame
    tf_b2.word_wrap = True
    p = tf_b2.paragraphs[0]
    p.text = "The Stemroom Approach"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_OK

    points_2 = [
        "A focused, distraction-free web platform tailored to audio workflows.",
        "Explicit project management with deadlines and collaborator rosters.",
        "Quality gatekeeping: Content moderation as a core architectural step.",
        "Persistent musician portfolio with strict status transparency.",
        "Robust, zero-bloat stack built with native Servlets, JDBC and MySQL."
    ]
    for pt in points_2:
        p_pt = tf_b2.add_paragraph()
        p_pt.text = "• " + pt
        p_pt.font.name = FONT_BODY
        p_pt.font.size = Pt(13)
        p_pt.font.color.rgb = INK_MUTED

    # =========================================================================
    # SLIDE 3: Two User Types and Capabilities
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "2. User Roles & Core Responsibilities")

    # Card 1: Musician Role
    create_card(s3, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), bg=PANEL_COLOR, border=COLOR_OK)
    m_box = s3.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.1), Inches(4.7))
    tf_m = m_box.text_frame
    tf_m.word_wrap = True

    pm0 = tf_m.paragraphs[0]
    pm0.text = "ROLE: MUSICIAN"
    pm0.font.name = FONT_HEADING
    pm0.font.size = Pt(12)
    pm0.font.bold = True
    pm0.font.color.rgb = COLOR_OK

    pm1 = tf_m.add_paragraph()
    pm1.text = "Creative Portfolio & Collaboration"
    pm1.font.name = FONT_HEADING
    pm1.font.size = Pt(18)
    pm1.font.bold = True
    pm1.font.color.rgb = INK

    m_features = [
        ("Project Lifecycles", "Create projects with title, description, and strict target deadlines."),
        ("Team Collaboration", "Invite fellow musicians by email; manage member permissions and status."),
        ("Portfolio Management", "Upload songs with lossless/compressed audio files (MP3, WAV, FLAC)."),
        ("Status Tracking", "Immediate visibility into song moderation states (Pending, Approved, Rejected)."),
        ("Role Boundaries", "Guarded by AuthFilter; strictly isolated from administrative endpoints.")
    ]
    for title, desc in m_features:
        p = tf_m.add_paragraph()
        p.text = f"• {title}: {desc}"
        p.font.name = FONT_BODY
        p.font.size = Pt(13)
        p.font.color.rgb = INK_MUTED

    # Card 2: Admin Role
    create_card(s3, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), bg=PANEL_COLOR, border=COLOR_WAIT)
    a_box = s3.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.7))
    tf_a = a_box.text_frame
    tf_a.word_wrap = True

    pa0 = tf_a.paragraphs[0]
    pa0.text = "ROLE: ADMINISTRATOR"
    pa0.font.name = FONT_HEADING
    pa0.font.size = Pt(12)
    pa0.font.bold = True
    pa0.font.color.rgb = COLOR_WAIT

    pa1 = tf_a.add_paragraph()
    pa1.text = "Platform Governance & Quality Control"
    pa1.font.name = FONT_HEADING
    pa1.font.size = Pt(18)
    pa1.font.bold = True
    pa1.font.color.rgb = INK

    a_features = [
        ("User Administration", "Provision accounts, toggle roles (Admin/Musician), and delete inactive users."),
        ("Cascading Safety", "User removal automatically and safely cascades to associated projects and songs."),
        ("Song Content Moderation", "Review submitted songs in a dedicated queue; approve or reject audio submissions."),
        ("System Configuration", "Toggle global public sign-ups on/off and enforce per-musician project limits."),
        ("Analytics Dashboard", "High-level overview of pending review queues and active user counts.")
    ]
    for title, desc in a_features:
        p = tf_a.add_paragraph()
        p.text = f"• {title}: {desc}"
        p.font.name = FONT_BODY
        p.font.size = Pt(13)
        p.font.color.rgb = INK_MUTED

    # =========================================================================
    # SLIDE 4: Scope for Review 1 vs Plan for Review 2
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "3. Scope: Review 1 Implementation vs Review 2 Roadmap")

    # Review 1 Implemented Card
    create_card(s4, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), bg=PANEL_COLOR, border=COLOR_OK)
    r1_box = s4.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.1), Inches(4.7))
    tf_r1 = r1_box.text_frame
    tf_r1.word_wrap = True

    p = tf_r1.paragraphs[0]
    p.text = "COMPLETED IN REVIEW 1 (WORKING NOW)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_OK

    p = tf_r1.add_paragraph()
    p.text = "Foundation & Operational Modules"
    p.font.name = FONT_HEADING
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = INK

    r1_items = [
        "Full project directory structure and Maven WAR deployment lifecycle.",
        "Relational database design (MySQL 8) with DDL, constraints & seed data.",
        "JDBC database layer with PreparedStatement and custom CrudDao generics.",
        "Session authentication & AuthFilter security boundary.",
        "Musician project collaboration (create, list, add members, update status).",
        "Song portfolio upload with multi-part disk storage.",
        "Admin user management, content moderation, and dynamic settings.",
        "Custom, responsive CSS design system with zero third-party frameworks."
    ]
    for item in r1_items:
        p = tf_r1.add_paragraph()
        p.text = "✔ " + item
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.font.color.rgb = INK_MUTED

    # Review 2 Planned Card
    create_card(s4, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), bg=PANEL_COLOR, border=LINE_COLOR)
    r2_box = s4.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.7))
    tf_r2 = r2_box.text_frame
    tf_r2.word_wrap = True

    p = tf_r2.paragraphs[0]
    p.text = "PLANNED FOR REVIEW 2"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_WAIT

    p = tf_r2.add_paragraph()
    p.text = "Advanced Audio & Real-time Collaboration"
    p.font.name = FONT_HEADING
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = INK

    r2_items = [
        "In-browser HTML5 audio streaming & waveform visualizer.",
        "Song editing workflow (metadata updates and audio file replacement).",
        "Project-level song sharing with track stems distribution.",
        "Time-stamped audio comments for precise musical feedback.",
        "Real-time notifications center and activity audit log.",
        "Enterprise auth enhancements (BCrypt/PBKDF2 hashing & CSRF tokens)."
    ]
    for item in r2_items:
        p = tf_r2.add_paragraph()
        p.text = "⏳ " + item
        p.font.name = FONT_BODY
        p.font.size = Pt(13)
        p.font.color.rgb = INK_MUTED

    # =========================================================================
    # SLIDE 5: Architecture Diagram & Request Flow
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "4. System Architecture & Request Lifecycle")

    # Pipeline Boxes: 5 Horizontal Steps
    steps = [
        ("1. Browser", "HTTP Request\n(GET / POST)\nForm Data / Audio", PANEL_COLOR, INK),
        ("2. AuthFilter", "Session verification\nRole check\nRoute protection", PANEL_COLOR, COLOR_WAIT),
        ("3. Servlet", "Input validation\nOrchestration\nView selection", PANEL_COLOR, ACCENT_MUSTARD),
        ("4. DAO Layer", "PreparedStatement\nRowMapper mapping\nConnection pool", PANEL_COLOR, COLOR_OK),
        ("5. MySQL 8", "Relational persistence\nACID integrity\nForeign key cascades", PANEL_COLOR, INK)
    ]

    card_width = Inches(2.15)
    gap = Inches(0.24)
    start_left = Inches(0.8)

    for i, (stitle, sdesc, sbg, scolor) in enumerate(steps):
        left_pos = start_left + i * (card_width + gap)
        create_card(s5, left_pos, Inches(1.8), card_width, Inches(2.3), bg=sbg, border=LINE_COLOR)

        tb = s5.shapes.add_textbox(left_pos + Inches(0.1), Inches(1.9), card_width - Inches(0.2), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = stitle
        p.font.name = FONT_HEADING
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = scolor

        p = tf.add_paragraph()
        p.text = sdesc
        p.font.name = FONT_BODY
        p.font.size = Pt(11)
        p.font.color.rgb = INK_MUTED

    # Downward Return Flow Card
    create_card(s5, Inches(0.8), Inches(4.4), Inches(11.733), Inches(2.5), bg=PANEL_COLOR, border=LINE_COLOR)
    ret_box = s5.shapes.add_textbox(Inches(1.1), Inches(4.5), Inches(11.1), Inches(2.3))
    tf_ret = ret_box.text_frame
    tf_ret.word_wrap = True

    p = tf_ret.paragraphs[0]
    p.text = "Return Pipeline: Servlet ➔ JSP View ➔ Escaped HTML Rendering"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = INK

    flow_points = [
        "JSP Isolation: Views reside strictly under /WEB-INF/views/ to prevent direct URL access from browsers.",
        "No Scriptlets: 100% pure JSTL (<c:forEach>, <c:if>, <c:out>) tags maintain separation of presentation from logic.",
        "Automatic XSS Prevention: All dynamic database text is output via <c:out value='...'/> to guarantee HTML escaping.",
        "Post/Redirect/Get Pattern: State-changing POST submissions redirect on success to prevent accidental duplicate submissions, while failures forward back with retained form inputs.",
        "No Heavy Framework Overhead: Built with clean Jakarta Servlet 6.0 and JDBC specifications to satisfy exact evaluation rubrics."
    ]
    for pt in flow_points:
        p = tf_ret.add_paragraph()
        p.text = "• " + pt
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.font.color.rgb = INK_MUTED

    # =========================================================================
    # SLIDE 6: Database Design & ER Diagram Decisions
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "5. Database Schema & Relational Design Decisions")

    # Table Schema Summary Card
    create_card(s6, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), bg=PANEL_COLOR, border=LINE_COLOR)
    sch_box = s6.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.1), Inches(4.7))
    tf_sch = sch_box.text_frame
    tf_sch.word_wrap = True

    p = tf_sch.paragraphs[0]
    p.text = "RELATIONAL SCHEMA ENTITIES"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_OK

    p = tf_sch.add_paragraph()
    p.text = "5 Core Normalized Tables in MySQL 8"
    p.font.name = FONT_HEADING
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = INK

    schema_info = [
        ("users", "id (PK), name, email (UK), password_hash, role (ENUM)"),
        ("projects", "id (PK), title, description, owner_id (FK), deadline, status (ENUM)"),
        ("project_members", "project_id (PK, FK), user_id (PK, FK) [Join Table]"),
        ("songs", "id (PK), owner_id (FK), title, album, file_name, status (ENUM)"),
        ("settings", "setting_key (PK), setting_value (VARCHAR)")
    ]
    for tbl, cols in schema_info:
        p = tf_sch.add_paragraph()
        p.text = f"• {tbl}: "
        p.font.name = FONT_HEADING
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = INK

        p_c = tf_sch.add_paragraph()
        p_c.text = f"  {cols}"
        p_c.font.name = FONT_BODY
        p_c.font.size = Pt(11)
        p_c.font.color.rgb = INK_MUTED

    # Design Decisions Card
    create_card(s6, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), bg=PANEL_COLOR, border=ACCENT_MUSTARD)
    dec_box = s6.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.7))
    tf_dec = dec_box.text_frame
    tf_dec.word_wrap = True

    p = tf_dec.paragraphs[0]
    p.text = "KEY ARCHITECTURAL DECISIONS"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_WAIT

    p = tf_dec.add_paragraph()
    p.text = "Data Integrity & Normalization"
    p.font.name = FONT_HEADING
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = INK

    decisions = [
        ("1. Single Owner Isolation", "Project ownership lives directly on `projects.owner_id`. Collaborator join table only stores invited peers, avoiding duplicate owner entries."),
        ("2. Explicit Moderation States", "Every uploaded track defaults to 'PENDING'. Only administrative approval unlocks full platform availability."),
        ("3. Dynamic Settings Store", "Key-value schema in `settings` allows toggling registrations and limits without executing database schema migrations."),
        ("4. Cascading Foreign Keys", "ON DELETE CASCADE guarantees that deleting an account completely purges associated projects and songs without orphaned rows."),
        ("5. Normalized Unique Emails", "Emails stored in strict lowercase with UNIQUE constraint to prevent case-variation duplicate accounts.")
    ]
    for d_title, d_desc in decisions:
        p = tf_dec.add_paragraph()
        p.text = d_title
        p.font.name = FONT_HEADING
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = INK

        p_sub = tf_dec.add_paragraph()
        p_sub.text = d_desc
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = INK_MUTED

    # =========================================================================
    # SLIDE 7: Core Java & JDBC Implementation
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "6. Core Java & Secure JDBC Implementation")

    # Left: Core Java Rubric Mapping
    create_card(s7, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), bg=PANEL_COLOR, border=LINE_COLOR)
    cj_box = s7.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.1), Inches(4.7))
    tf_cj = cj_box.text_frame
    tf_cj.word_wrap = True

    p = tf_cj.paragraphs[0]
    p.text = "OOP & LANGUAGE FEATURES"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_OK

    p = tf_cj.add_paragraph()
    p.text = "Demonstrated Syllabus Concepts"
    p.font.name = FONT_HEADING
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = INK

    java_features = [
        ("Inheritance", "Abstract User base class extended by Admin and Musician subclasses."),
        ("Polymorphism", "Overridden canModerate() and getRole() called uniformly without type checks."),
        ("Interfaces & Generics", "CrudDao<T> implemented across UserDao, ProjectDao, and SongDao."),
        ("Collections & Optionals", "List<T> for multi-row queries; Optional<T> ensures null is never returned."),
        ("Enums with Metadata", "Role, ProjectStatus, and ModerationStatus with parse() and labels."),
        ("Exception Wrapping", "DataAccessException cleanly wraps checked SQLExceptions."),
        ("Try-with-resources", "Guarantees connections, statements, and result sets close safely.")
    ]
    for feat, expl in java_features:
        p = tf_cj.add_paragraph()
        p.text = f"• {feat}: {expl}"
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.font.color.rgb = INK_MUTED

    # Right: JDBC Code Sample Card
    create_card(s7, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), bg=PANEL_COLOR, border=LINE_COLOR)
    code_box = s7.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.7))
    tf_code = code_box.text_frame
    tf_code.word_wrap = True

    p = tf_code.paragraphs[0]
    p.text = "SECURE JDBC CODE PATTERN"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_WAIT

    p = tf_code.add_paragraph()
    p.text = "Parameterized Queries & RowMapper"
    p.font.name = FONT_HEADING
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = INK

    sample_code = (
        "// Example from SongDao.java\n"
        "public Optional<Song> findById(int id) {\n"
        "    String sql = \"\"\"\n"
        "        SELECT s.*, u.name AS owner_name\n"
        "        FROM songs s\n"
        "        JOIN users u ON s.owner_id = u.id\n"
        "        WHERE s.id = ?\n"
        "        \"\"\";\n"
        "    return Db.queryOne(sql, this::mapRow, id);\n"
        "}\n\n"
        "private Song mapRow(ResultSet rs) throws SQLException {\n"
        "    return new Song(\n"
        "        rs.getInt(\"id\"),\n"
        "        rs.getInt(\"owner_id\"),\n"
        "        rs.getString(\"title\"),\n"
        "        rs.getString(\"album\"),\n"
        "        rs.getString(\"file_name\"),\n"
        "        ModerationStatus.parse(rs.getString(\"status\"))\n"
        "    );\n"
        "}"
    )
    p = tf_code.add_paragraph()
    p.text = sample_code
    p.font.name = "Courier New"
    p.font.size = Pt(10)
    p.font.color.rgb = INK

    # =========================================================================
    # SLIDE 8: Application Screens & UI Showcase
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "7. Live Application Interface & Visual Validation")

    # 4 Screen Quadrants
    screens = [
        ("Musician Dashboard", "docs/screenshots/02-musician-dashboard.png", Inches(0.8), Inches(1.8)),
        ("Project Collaboration", "docs/screenshots/04-project-detail.png", Inches(6.8), Inches(1.8)),
        ("Portfolio & Uploads", "docs/screenshots/05-portfolio.png", Inches(0.8), Inches(4.5)),
        ("Admin Moderation Queue", "docs/screenshots/07-admin-moderation.png", Inches(6.8), Inches(4.5))
    ]

    for label, img_path, left, top in screens:
        create_card(s8, left, top, Inches(5.7), Inches(2.5), bg=PANEL_COLOR, border=LINE_COLOR)

        # Label
        lbl_box = s8.shapes.add_textbox(left + Inches(0.1), top + Inches(0.05), Inches(5.5), Inches(0.3))
        p = lbl_box.text_frame.paragraphs[0]
        p.text = label.upper()
        p.font.name = FONT_HEADING
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_WAIT

        if os.path.exists(img_path):
            try:
                s8.shapes.add_picture(img_path, left + Inches(0.2), top + Inches(0.4), width=Inches(5.3), height=Inches(2.0))
            except Exception:
                pass

    # =========================================================================
    # SLIDE 9: Design System & Intentional Aesthetics
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "8. Design System: Aesthetic Choices & Refused Cliches")

    # Left: Philosophy & Color Tokens
    create_card(s9, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), bg=PANEL_COLOR, border=COLOR_OK)
    col_box = s9.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.1), Inches(4.7))
    tf_col = col_box.text_frame
    tf_col.word_wrap = True

    p = tf_col.paragraphs[0]
    p.text = "PHILOSOPHY & PALETTE"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_OK

    p = tf_col.add_paragraph()
    p.text = '"A rehearsal room with the lights on"'
    p.font.name = FONT_HEADING
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = INK

    palette_tokens = [
        ("Studio Canvas (#e5e8e2)", "Pale grey-green page background, soothing for late-night editing."),
        ("Deep Ink (#1c2220)", "High-contrast green-black text and structural borders."),
        ("Mustard Fill (#e0a526)", "Reserved exclusively for primary action buttons; never as decoration."),
        ("Action Crimson (#a8362a)", "Used solely for destructive deletes and error boundaries."),
        ("Approved Green (#2c6a43)", "Verification tags with paired square icons for full greyscale accessibility."),
        ("Typography Pairing", "Franklin Gothic condensed headings paired with legible, wide Verdana body text.")
    ]
    for name, expl in palette_tokens:
        p = tf_col.add_paragraph()
        p.text = f"• {name}: {expl}"
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.font.color.rgb = INK_MUTED

    # Right: The Refused Defaults (from docs/DESIGN.md)
    create_card(s9, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), bg=PANEL_COLOR, border=COLOR_DANGER)
    ref_box = s9.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.7))
    tf_ref = ref_box.text_frame
    tf_ref.word_wrap = True

    p = tf_ref.paragraphs[0]
    p.text = "INTENTIONAL RESTRAINTS"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_DANGER

    p = tf_ref.add_paragraph()
    p.text = "What Was Deliberately Left Out"
    p.font.name = FONT_HEADING
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = INK

    refused_cliches = [
        ("No Generic AI Warm Cream & Terracotta", "Avoided cookie-cutter blog aesthetic."),
        ("No Neon / Acid-Green on Pitch Black", "Avoided clichéd 'cyberpunk' developer tropes."),
        ("No Rounded Drop-Shadow Cards", "Information layout follows editorial structure, not SaaS card grids."),
        ("No Monospace Labels for Plain Data", "Keeps technical metadata natural to read for musicians."),
        ("No Cluttered 'Welcome Back' Hero Banners", "Musicians immediately see actionable projects and tracks."),
        ("Zero CSS Framework Bloat", "Handcrafted, lightweight CSS (~400 lines) with rapid sub-10ms paint times.")
    ]
    for r_title, r_desc in refused_cliches:
        p = tf_ref.add_paragraph()
        p.text = f"✕ {r_title}: "
        p.font.name = FONT_HEADING
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_DANGER

        p_desc = tf_ref.add_paragraph()
        p_desc.text = f"   {r_desc}"
        p_desc.font.name = FONT_BODY
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = INK_MUTED

    # =========================================================================
    # SLIDE 10: Conclusion & Review 2 Roadmap
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "9. Review 1 Conclusion & Demonstration")

    # Card 1: Deliverables Delivered
    create_card(s10, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), bg=PANEL_COLOR, border=COLOR_OK)
    c1_box = s10.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.1), Inches(4.7))
    tf_c1 = c1_box.text_frame
    tf_c1.word_wrap = True

    p = tf_c1.paragraphs[0]
    p.text = "REVIEW 1 STATUS"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_OK

    p = tf_c1.add_paragraph()
    p.text = "All Rubric Goals Satisfied"
    p.font.name = FONT_HEADING
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = INK

    deliverables = [
        ("Problem Scope", "Online Music Collaboration Platform (Problem Statement 22) fully addressed."),
        ("Database Architecture", "MySQL 8 database normalized with cascading constraints & seed scripts."),
        ("Core Java Mastery", "Inheritance, Polymorphism, Generics, Collections, and Custom Exceptions."),
        ("Web Layer", "Jakarta Servlet 6.0, JSTL 3.0, and secure session AuthFilter."),
        ("Packaging & Deployment", "Apache Tomcat 10.1 WAR deployment verified and running."),
        ("Open Source Delivery", "Version controlled on GitHub (omsingh02/stemroom-byteforce).")
    ]
    for d_h, d_b in deliverables:
        p = tf_c1.add_paragraph()
        p.text = f"✔ {d_h}: {d_b}"
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.font.color.rgb = INK_MUTED

    # Card 2: Next Steps & Q&A
    create_card(s10, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), bg=PANEL_COLOR, border=ACCENT_MUSTARD)
    c2_box = s10.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.7))
    tf_c2 = c2_box.text_frame
    tf_c2.word_wrap = True

    p = tf_c2.paragraphs[0]
    p.text = "PROJECT ROADMAP & DEMO"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_WAIT

    p = tf_c2.add_paragraph()
    p.text = "Ready for Evaluation"
    p.font.name = FONT_HEADING
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = INK

    p = tf_c2.add_paragraph()
    p.text = "\nLive Demo Accounts Ready on http://localhost:8080/stemroom/:\n" \
             "• Admin: admin@stemroom.local / admin123\n" \
             "• Musician (Mira): mira@stemroom.local / musician123\n" \
             "• Musician (Dev): dev@stemroom.local / musician123\n"
    p.font.name = FONT_BODY
    p.font.size = Pt(13)
    p.font.color.rgb = INK

    p = tf_c2.add_paragraph()
    p.text = "Repository:\nhttps://github.com/omsingh02/stemroom-byteforce"
    p.font.name = FONT_BODY
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_OK

    p = tf_c2.add_paragraph()
    p.text = "\nThank you!\nQuestions & Feedback Welcome."
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = INK

    # Save output
    output_path = "Stemroom_Review_1_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    build_presentation()
